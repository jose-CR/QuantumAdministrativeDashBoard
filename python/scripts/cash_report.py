import json
import sys
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
MONEY = "#,##0.00"
DATE_FMT = "yyyy-mm-dd"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")


# ---------- generic helpers ----------
def put(ws, row, col, value=None, *, bold=False, color=None, size=10, fmt=None, fill=None, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def clean(v):
    if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date()
    if isinstance(v, np.generic):
        return v.item()
    return v


def load_df(headers, rows, cfg):
    df = pd.DataFrame(rows, columns=headers)
    df[cfg["amount"]] = pd.to_numeric(df[cfg["amount"]], errors="coerce")
    df[cfg["date"]] = pd.to_datetime(df[cfg["date"]], errors="coerce")
    g = cfg.get("group_by")
    if g:
        df[g] = df[g].where(df[g].notna() & (df[g].astype(str).str.strip() != ""), "Sin dato")
    return df


def find_column(ws, header, header_row=1):
    for c in ws[header_row]:
        if str(c.value or "").strip().lower() == str(header).strip().lower():
            return c.column
    raise ValueError(f"Column '{header}' not found in row {header_row}")


# ---------- duplicate / anomaly detection (works on any df + cfg) ----------
def analyze(cfg, df, today):
    """Return {row position (0-indexed within df): [reasons]}."""
    found = {}

    def flag(mask, reason):
        for i in df.index[mask]:
            found.setdefault(int(i), []).append(reason)

    amt, dt = cfg["amount"], cfg["date"]
    inv, party = cfg.get("invoice"), cfg.get("party")

    key = df.copy()
    for c in (inv, party):
        if c:
            key[c] = key[c].where(key[c].astype(str).str.strip() != "", np.nan)

    exact = pd.Series(False, index=df.index)
    if inv and party:
        exact = key.duplicated([inv, party, amt, dt], keep=False) & key[inv].notna()
        flag(exact, "Duplicado exacto (factura, contraparte, monto y fecha)")
        flag(
            key.duplicated([inv, party], keep=False) & key[inv].notna() & ~exact,
            "Misma factura y contraparte con datos distintos (¿pago parcial o error?)",
        )
    if party:
        flag(
            key.duplicated([party, amt, dt], keep=False)
            & key[party].notna() & key[amt].notna() & key[dt].notna() & ~exact,
            "Mismo monto, fecha y contraparte (posible doble registro)",
        )

    s = df[amt]
    flag(s.isna(), "Monto vacío o no numérico")
    flag(s <= 0, "Monto cero o negativo")
    valid = s.dropna()
    if len(valid) >= 8:
        q1, q3 = valid.quantile([0.25, 0.75])
        iqr = q3 - q1
        hi, lo = q3 + 1.5 * iqr, q1 - 1.5 * iqr
        flag(s > hi, f"Monto atípico alto (> {hi:,.2f})")
        flag((s < lo) & (s > 0), f"Monto atípico bajo (< {lo:,.2f})")

    d = df[dt]
    flag(d.isna(), "Fecha vacía o inválida")
    flag(d > pd.Timestamp(today), "Fecha futura")
    return found


def rng(col, flow):
    return f"'{flow['sheet']}'!${L(col)}${flow['start']}:${L(col)}${flow['last']}"


# ---------- write/annotate the data sheet ----------
def build_report_sheet(wb, cfgs, dfs, flags, gap):
    """Combined mode: write both flows side by side on a new sheet ('Reporte')."""
    ws = wb.active
    ws.title = "Reporte"
    data_row = 3  # row 1: title, row 2: headers, data from row 3
    col0 = 1
    flows = []

    for cfg, df, fl in zip(cfgs, dfs, flags):
        headers = cfg["headers"]
        put(ws, 1, col0, cfg["title"].upper(), bold=True, size=14)
        for j, h in enumerate(headers):
            put(ws, 2, col0 + j, h, bold=True, color="FFFFFF", fill=HEAD_FILL,
                align=Alignment(horizontal="center", vertical="center", wrap_text=True))

        a_idx = headers.index(cfg["amount"])
        d_idx = headers.index(cfg["date"])
        for i, rec in enumerate(df.itertuples(index=False, name=None)):
            for j, v in enumerate(rec):
                put(ws, data_row + i, col0 + j, clean(v),
                    fmt=MONEY if j == a_idx else DATE_FMT if j == d_idx else None,
                    fill=FLAG_FILL if i in fl else None)

        last = max(data_row + len(df) - 1, data_row)
        total_row = last + 2
        a = col0 + a_idx
        if a_idx > 0:
            put(ws, total_row, a - 1, "TOTAL", bold=True)
        put(ws, total_row, a, f"=SUM({L(a)}{data_row}:{L(a)}{last})", bold=True, fmt=MONEY)

        for j, h in enumerate(headers):
            if j == a_idx:
                w = 14
            elif j == d_idx:
                w = 12
            else:
                w = min(max([len(h)] + [len(str(v)) for v in df.iloc[:, j].dropna()] + [10]) + 2, 45)
            ws.column_dimensions[L(col0 + j)].width = w

        g = cfg.get("group_by")
        flows.append({
            "title": cfg["title"], "sheet": "Reporte", "start": data_row, "last": last,
            "a": a, "d": col0 + d_idx, "g": col0 + headers.index(g) if g else None,
        })
        col0 += len(headers) + gap

    for k in range(len(cfgs[0]["headers"]) + 1, len(cfgs[0]["headers"]) + gap + 1):
        ws.column_dimensions[L(k)].width = 4
    ws.freeze_panes = "A3"
    ws.row_dimensions[2].height = 30
    return flows


def annotate_existing_sheet(ws, cfg, df, fl):
    """Single-flow mode: the sheet already has data (written by Laravel/Filament).
    Give it the same look as the combined report (title row + styled header),
    highlight flagged rows, cast the amount/date columns, and add the TOTAL row."""
    orig_headers = [str(c.value) for c in ws[1]]

    # push the existing header+data down one row and add a title row, matching
    # the "Reporte" sheet layout: row 1 = title, row 2 = header, data from row 3
    ws.insert_rows(1)
    header_row = 2
    put(ws, 1, 1, cfg["title"].upper(), bold=True, size=14)
    for c in ws[header_row]:
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30
    ws.freeze_panes = f"A{header_row + 1}"

    a_col = find_column(ws, cfg["amount"], header_row)
    d_col = find_column(ws, cfg["date"], header_row)
    start = header_row + 1
    last = ws.max_row

    for i in range(len(df)):
        row = start + i
        for c in range(1, len(orig_headers) + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(name=FONT, size=10)
            if i in fl:
                cell.fill = FLAG_FILL

        cell = ws.cell(row=row, column=a_col)
        if cell.value not in (None, ""):
            cell.value = float(cell.value)
        cell.number_format = MONEY

        # the sheet was written by Laravel/OpenSpout as plain CSV text, so the date
        # column is a string cell; replace it with a real date so SUMIFS/EDATE work
        dcell = ws.cell(row=row, column=d_col)
        parsed = df[cfg["date"]].iloc[i]
        dcell.value = parsed.to_pydatetime().date() if pd.notna(parsed) else None
        dcell.number_format = DATE_FMT

    total_row = last + 2
    if a_col > 1:
        put(ws, total_row, a_col - 1, "TOTAL", bold=True)
    letter = L(a_col)
    put(ws, total_row, a_col, f"=SUM({letter}{start}:{letter}{last})", bold=True, fmt=MONEY)

    for j, h in enumerate(orig_headers):
        col = j + 1
        if col == a_col:
            w = 14
        elif col == d_col:
            w = 12
        else:
            w = min(max([len(h)] + [len(str(v)) for v in df.iloc[:, j].dropna().astype(str)] + [10]) + 2, 45)
        ws.column_dimensions[L(col)].width = w

    g = cfg.get("group_by")
    return {
        "title": cfg["title"], "sheet": ws.title, "start": start, "last": last,
        "a": a_col,
        "d": d_col,
        "g": find_column(ws, g, header_row) if g else None,
    }


# ---------- sheet "Resumen" ----------
def write_summary(wb, flows, dfs, cfgs):
    ws = wb.create_sheet("Resumen")
    put(ws, 1, 1, "RESUMEN MENSUAL", bold=True, size=14)

    single = len(flows) == 1
    amount_cols = ["Monto"] if single else ["Entradas", "Salidas"]
    headers = ["Mes"] + amount_cols + (["Neto"] if not single else []) + ["Inicio de mes"]
    for j, h in enumerate(headers, start=1):
        put(ws, 2, j, h, bold=True, color="FFFFFF", fill=HEAD_FILL, align=Alignment(horizontal="center"))
    date_col = len(headers)

    all_dates = pd.concat([d[c["date"]] for c, d in zip(cfgs, dfs)]).dropna()
    today = date.today()
    all_dates = all_dates[all_dates <= pd.Timestamp(today)]
    periods = list(pd.period_range(all_dates.min().to_period("M"), all_dates.max().to_period("M"), freq="M")) \
        if len(all_dates) else []

    r = 3
    first = r
    for p in periods:
        put(ws, r, 1, str(p))
        put(ws, r, date_col, p.start_time.date(), fmt=DATE_FMT, color="808080")
        for j, flow in enumerate(flows, start=2):
            put(ws, r, j,
                f'=SUMIFS({rng(flow["a"], flow)},{rng(flow["d"], flow)},">="&${L(date_col)}{r},'
                f'{rng(flow["d"], flow)},"<"&EDATE(${L(date_col)}{r},1))', fmt=MONEY)
        if not single:
            put(ws, r, 4, f"=B{r}-C{r}", fmt=MONEY)
        r += 1
    last = r - 1
    if periods:
        put(ws, r, 1, "TOTAL", bold=True)
        for j in range(2, date_col):
            put(ws, r, j, f"=SUM({L(j)}{first}:{L(j)}{last})", bold=True, fmt=MONEY)
        r += 1
        put(ws, r, 1, "Excluye registros sin fecha válida o con fecha futura.", color="808080")
        r += 1

    widths = [14] + [16] * (len(amount_cols) + (0 if single else 1)) + [14]
    for c, w in zip([L(i) for i in range(1, date_col + 1)], widths):
        ws.column_dimensions[c].width = w

    if periods:
        bar = BarChart()
        bar.type, bar.grouping = "col", "clustered"
        bar.title = "Monto por mes" if single else "Entradas vs Salidas por mes"
        bar.add_data(Reference(ws, min_col=2, max_col=1 + len(amount_cols), min_row=2, max_row=last),
                     titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        bar.y_axis.title = "Monto"
        bar.x_axis.delete = bar.y_axis.delete = False
        bar.width, bar.height = 18, 8
        ws.add_chart(bar, "G2")

    r += 2
    anchors = ["G19", "O19"]
    for cfg, df, flow, anchor in zip(cfgs, dfs, flows, anchors):
        g = cfg.get("group_by")
        if not g:
            continue
        put(ws, r, 1, f'{cfg["title"].upper()} POR {g.upper()}', bold=True)
        r += 1
        put(ws, r, 1, g, bold=True, color="FFFFFF", fill=HEAD_FILL)
        put(ws, r, 2, "Monto", bold=True, color="FFFFFF", fill=HEAD_FILL)
        hdr = r
        groups = df.groupby(g)[cfg["amount"]].sum().sort_values(ascending=False).index
        for k in groups:
            r += 1
            put(ws, r, 1, str(k))
            put(ws, r, 2, f'=SUMIFS({rng(flow["a"], flow)},{rng(flow["g"], flow)},A{r})', fmt=MONEY)
        if len(groups):
            pie = PieChart()
            pie.title = f'{cfg["title"]} por {g}'
            pie.add_data(Reference(ws, min_col=2, min_row=hdr, max_row=r), titles_from_data=True)
            pie.set_categories(Reference(ws, min_col=1, min_row=hdr + 1, max_row=r))
            pie.dataLabels = DataLabelList()
            pie.dataLabels.showPercent = True
            pie.dataLabels.showVal = pie.dataLabels.showCatName = pie.dataLabels.showSerName = False
            pie.dataLabels.showLeaderLines = False
            pie.width, pie.height = 12, 8
            ws.add_chart(pie, anchor)
        r += 3
    return periods


# ---------- sheet "Análisis" ----------
def write_analysis(wb, cfgs, dfs, flags, flows, periods, horizon):
    ws = wb.create_sheet("Análisis")
    today = date.today()
    single = len(flows) == 1

    put(ws, 1, 1, "ANÁLISIS: DUPLICADOS, ANOMALÍAS Y PROYECCIÓN", bold=True, size=14)
    put(ws, 3, 1, "1. Duplicados y anomalías (filas resaltadas en la hoja de datos)", bold=True, size=11)
    heads = (["Hoja"] if not single else []) + ["Fila", "Factura", "Contraparte", "Fecha", "Monto", "Motivo"]
    for j, h in enumerate(heads, start=1):
        put(ws, 4, j, h, bold=True, color="FFFFFF", fill=HEAD_FILL, align=Alignment(horizontal="center"))

    r = 5
    for cfg, df, fl, flow in zip(cfgs, dfs, flags, flows):
        for pos in sorted(fl):
            row = df.iloc[pos]
            base = [flow["title"]] if not single else []
            vals = base + [
                flow["start"] + pos,
                clean(row[cfg["invoice"]]) if cfg.get("invoice") else None,
                clean(row[cfg["party"]]) if cfg.get("party") else None,
                clean(row[cfg["date"]]), clean(row[cfg["amount"]]), "; ".join(fl[pos]),
            ]
            for j, v in enumerate(vals, start=1):
                put(ws, r, j, v, fmt=MONEY if heads[j - 1] == "Monto" else DATE_FMT if heads[j - 1] == "Fecha" else None)
            r += 1
    if r == 5:
        put(ws, 5, 1, "Sin hallazgos")
        r = 6

    put(ws, r + 1, 1, "Reglas: duplicado exacto = misma factura, contraparte, monto y fecha; atípico = fuera de "
                      "Q1-1.5·IQR / Q3+1.5·IQR (solo con 8 o más registros); fecha futura = posterior a hoy. "
                      "Una misma factura con datos distintos puede ser un pago parcial legítimo.", color="808080")
    r += 4

    put(ws, r, 1, "2. Proyección a partir del histórico", bold=True, size=11)
    r += 1
    amount_cols = ["Monto"] if single else ["Entradas", "Salidas"]
    if not periods:
        put(ws, r, 1, "Sin datos con fecha para proyectar.")
    else:
        in_progress = periods[-1] == pd.Period(today, "M")
        complete = periods[:-1] if in_progress else periods
        heads2 = ["Mes"] + amount_cols + (["Neto"] if not single else []) + ["Tipo", "n"]
        for j, h in enumerate(heads2, start=1):
            put(ws, r, j, h, bold=True, color="FFFFFF", fill=HEAD_FILL, align=Alignment(horizontal="center"))
        hdr = r
        neto_col = len(amount_cols) + 2
        tipo_col = neto_col + (1 if not single else 0)
        n_col = tipo_col + 1
        for i, p in enumerate(periods):
            r += 1
            src = 3 + i
            put(ws, r, 1, str(p))
            for j in range(len(amount_cols)):
                put(ws, r, 2 + j, f"=Resumen!{L(2 + j)}{src}", fmt=MONEY)
            if not single:
                put(ws, r, neto_col, f"=B{r}-C{r}", fmt=MONEY)
            put(ws, r, tipo_col, "Real (mes en curso, parcial)" if in_progress and i == len(periods) - 1 else "Real")
            put(ws, r, n_col, i + 1)
        h0, h1 = hdr + 1, hdr + len(complete)
        if len(complete) >= 3:
            for k in range(1, horizon + 1):
                r += 1
                put(ws, r, 1, f"{periods[-1] + k} (proy.)")
                for j in range(len(amount_cols)):
                    c = L(2 + j)
                    put(ws, r, 2 + j, f"=MAX(0,FORECAST(${L(n_col)}{r},{c}${h0}:{c}${h1},$F${h0}:$F${h1}))"
                        .replace("$F$", f"${L(n_col)}$"), fmt=MONEY)
                if not single:
                    put(ws, r, neto_col, f"=B{r}-C{r}", fmt=MONEY)
                put(ws, r, tipo_col, "Proyección")
                put(ws, r, n_col, len(periods) + k)
            end = r
            r += 2
            put(ws, r, 1, "Prom. 3 meses completos", bold=True)
            for j in range(len(amount_cols)):
                c = L(2 + j)
                put(ws, r, 2 + j, f"=AVERAGE({c}{h1 - 2}:{c}{h1})", bold=True, fmt=MONEY)
            r += 1
            put(ws, r, 1, "Método: regresión lineal (FORECAST) sobre los meses completos; no considera "
                          "estacionalidad. Es una referencia orientativa: con pocos meses de historial la "
                          "incertidumbre es alta.", color="808080")

            line = LineChart()
            line.title = "Monto real y proyección" if single else "Entradas y salidas: real y proyección"
            line.add_data(Reference(ws, min_col=2, max_col=1 + len(amount_cols), min_row=hdr, max_row=end),
                          titles_from_data=True)
            line.set_categories(Reference(ws, min_col=1, min_row=hdr + 1, max_row=end))
            line.x_axis.delete = line.y_axis.delete = False
            for serie in line.series:
                serie.smooth = False
            line.width, line.height = 20, 9
            ws.add_chart(line, f"A{r + 3}")
        else:
            r += 2
            put(ws, r, 1, "Historial insuficiente: se necesitan al menos 3 meses completos para proyectar.")

    for c, w in zip("ABCDEFGH", (26, 16, 18, 28, 26, 14, 90, 20)):
        ws.column_dimensions[c].width = w


# ---------- mode: only append a TOTAL row (no analysis sheets) ----------
def add_total_only(payload):
    path = payload["path"]
    columns = payload.get("columns") or (
        [payload["amount_header"]] if payload.get("amount_header") else ["Monto", "Amount"]
    )
    wanted = {str(c).strip().lower() for c in columns}

    wb = load_workbook(path)
    ws = wb.active
    targets = [c.column for c in ws[1] if str(c.value or "").strip().lower() in wanted]
    if not targets:
        raise ValueError(f"No column matching {columns} found in row 1")

    last = ws.max_row
    total_row = last + 2
    totals = {}
    for col in targets:
        total = 0.0
        for r in range(2, last + 1):
            cell = ws.cell(row=r, column=col)
            if cell.value in (None, ""):
                continue
            cell.value = float(cell.value)
            cell.number_format = MONEY
            total += cell.value
        letter = L(col)
        value = ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{max(last, 2)})")
        value.number_format = MONEY
        value.font = Font(bold=True, name=FONT)
        totals[str(ws.cell(row=1, column=col).value)] = round(total, 2)
    first = targets[0]
    if first > 1:
        ws.cell(row=total_row, column=first - 1, value="TOTAL").font = Font(bold=True, name=FONT)
    wb.save(path)
    return {"path": path, "rows": last - 1, "totals": totals}


# ---------- mode: full report for one flow, appended to an existing sheet ----------
def single_flow_report(payload):
    path = payload["path"]
    cfg = {k: payload[k] for k in ("title", "amount", "date") }
    for k in ("invoice", "party", "group_by"):
        if payload.get(k):
            cfg[k] = payload[k]

    wb = load_workbook(path)
    ws = wb.active
    headers = [str(c.value) for c in ws[1]]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    df = load_df(headers, rows, cfg)
    today = date.today()
    fl = analyze(cfg, df, today)

    flow = annotate_existing_sheet(ws, cfg, df, fl)
    periods = write_summary(wb, [flow], [df], [cfg])
    write_analysis(wb, [cfg], [df], [fl], [flow], periods, int(payload.get("projection_months", 3)))

    wb.save(path)
    return {
        "path": path, "rows": len(df),
        "total": round(float(df[cfg["amount"]].sum()), 2),
        "flagged": len(fl),
    }


# ---------- mode: full report combining both flows into a new file ----------
def combined_report(p):
    cfgs = [p["inflows"], p["outflows"]]
    dfs = [load_df(c["headers"], c["rows"], c) for c in cfgs]
    today = date.today()
    flags = [analyze(c, d, today) for c, d in zip(cfgs, dfs)]

    wb = Workbook()
    flows = build_report_sheet(wb, cfgs, dfs, flags, int(p.get("gap", 2)))
    periods = write_summary(wb, flows, dfs, cfgs)
    write_analysis(wb, cfgs, dfs, flags, flows, periods, int(p.get("projection_months", 3)))
    wb.save(p["out"])

    return {
        "path": p["out"],
        "inflows": {"rows": len(dfs[0]), "total": round(float(dfs[0][cfgs[0]["amount"]].sum()), 2)},
        "outflows": {"rows": len(dfs[1]), "total": round(float(dfs[1][cfgs[1]["amount"]].sum()), 2)},
        "flagged": {c["title"]: len(f) for c, f in zip(cfgs, flags)},
    }


def main():
    p = json.load(sys.stdin)
    if "inflows" in p:
        result = combined_report(p)
    elif "amount" in p and "date" in p:
        result = single_flow_report(p)
    else:
        result = add_total_only(p)
    json.dump(result, sys.stdout)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(e, file=sys.stderr)
        sys.exit(1)